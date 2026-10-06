"""GIOŚ archive sheets → bronze.pm_hourly: header detection, number parsing, time conversion,
idempotency. Test workbooks mimic the real layout (6 header rows, wide format)."""

import datetime as dt
import gzip
import zipfile
from pathlib import Path

import openpyxl
import pytest

from smogcast.bronze.pm_archive import sheet_to_long_csv
from smogcast.bronze.steps import step_pm_hourly, step_station_meta


def make_sheet(path: Path, pollutant: str, positions: list[str], rows: list[tuple]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    n = len(positions)
    ws.append(["Nr", *range(1, n + 1)])
    ws.append(["Kod stacji", *[p.split("-")[0] for p in positions]])
    ws.append(["Wskaźnik", *[pollutant] * n])
    ws.append(["Czas uśredniania", *["1g"] * n])
    ws.append(["Jednostka", *["ug/m3"] * n])
    ws.append(["Kod stanowiska", *positions])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


ROWS = [
    (dt.datetime(2024, 1, 1, 1, 0), 12.5, None),
    (dt.datetime(2024, 1, 1, 2, 0, 0, 5000), "30,25", 7),       # ms noise, decimal comma
    (dt.datetime(2024, 1, 1, 2, 59, 59, 990000), "n/a", 8.0),   # noise below the hour, junk text
]


def test_sheet_to_long_csv(tmp_path):
    xlsx = make_sheet(tmp_path / "s.xlsx", "PM10", ["MpKrakAlKras-PM10-1g", "PmGdaLeczk08-PM10-1g"], ROWS)
    stats = sheet_to_long_csv(xlsx, tmp_path / "out.csv.gz")
    assert (stats.positions, stats.hours, stats.values, stats.empty, stats.unparseable) == (2, 3, 4, 1, 1)
    with gzip.open(tmp_path / "out.csv.gz", "rt", encoding="utf-8") as f:
        lines = f.read().splitlines()
    assert lines[0] == "position_code,time_cet_end,value"
    assert "MpKrakAlKras-PM10-1g,2024-01-01 02:00:00,30.25" in lines
    assert "PmGdaLeczk08-PM10-1g,2024-01-01 03:00:00,8.0" in lines   # 02:59:59.99 rounded up


def test_missing_header_row_is_an_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["something else"])
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError, match="Kod stanowiska"):
        sheet_to_long_csv(tmp_path / "bad.xlsx", tmp_path / "bad.csv.gz")


def _year_zip(ctx, year: int) -> None:
    raw = Path(ctx.storage.landing("gios_archive", "raw"))
    raw.mkdir(parents=True, exist_ok=True)
    tmp = raw / "tmp"
    tmp.mkdir()
    make_sheet(tmp / f"{year}_PM10_1g.xlsx", "PM10", ["MpKrakAlKras-PM10-1g"], ROWS)
    make_sheet(tmp / f"{year}_PM25_1g.xlsx", "PM2.5", ["MpKrakAlKras-PM2.5-1g"], ROWS)
    with zipfile.ZipFile(raw / f"{year}.zip", "w") as z:
        for f in tmp.iterdir():
            z.write(f, f.name)


def test_step_pm_hourly_converts_time_and_is_idempotent(ctx):
    ctx.cfg["run"]["archive_years"] = [2024]
    _year_zip(ctx, 2024)
    step_pm_hourly(ctx)
    step_pm_hourly(ctx)  # re-run must not duplicate rows

    df = ctx.storage.read("bronze", "pm_hourly")
    assert df.count() == 4  # per pollutant: 3 hours minus the junk "n/a" cell = 2 values
    assert {r["pollutant"] for r in df.select("pollutant").distinct().collect()} == {"PM10", "PM2.5"}
    first = df.where("pollutant = 'PM10'").orderBy("time_utc").first()
    # "2024-01-01 01:00 CET" (end of hour) = hour starting 2023-12-31 23:00 UTC
    assert first["time_utc"] == dt.datetime(2023, 12, 31, 23)
    assert first["station_code"] == "MpKrakAlKras" and first["source_year"] == 2024


def test_step_station_meta(ctx):
    raw = Path(ctx.storage.landing("gios_archive", "raw"))
    raw.mkdir(parents=True)
    wb = openpyxl.Workbook()
    st = wb.active
    st.title = "STACJE"
    st.append(["Nr", "Kod stacji", "Kod międzynarodowy", "Nazwa stacji", "Stary Kod stacji \n(o ile inny od aktualnego)",
               "Data uruchomienia", "Data zamknięcia", "Typ stacji", "Typ obszaru", "Rodzaj stacji", "Województwo",
               "Miejscowość", "Adres", "WGS84 φ N", "WGS84 λ E"])
    st.append([1, "MpKrakAlKras", "PL0012A", "Kraków, Aleja Krasińskiego", "MpKrakowWIOSAKra6117", dt.datetime(2003, 1, 1),
               None, "komunikacyjna", "miejski", "kontenerowa stacjonarna", "MAŁOPOLSKIE", "Kraków", "al. Krasińskiego",
               "50.057678", "19.926189"])
    pos = wb.create_sheet("STANOWISKA")
    pos.append(["Nr", "Kod stanowiska", "Kod stacji", "Nazwa stacji", "Wskaźnik - kod", "Wskaźnik", "Czas uśredniania",
                "Typ pomiaru", "Data uruchomienia", "Data zamknięcia", "Województwo"])
    pos.append([1, "MpKrakAlKras-PM10-1g", "MpKrakAlKras", "Kraków", "PM10", "pył", "1-godzinny", "automatyczny",
                dt.datetime(2003, 1, 1), None, "MAŁOPOLSKIE"])
    wb.save(raw / "metadata.xlsx")

    step_station_meta(ctx)
    s = ctx.storage.read("bronze", "gios_stations").first()
    assert (s["station_code"], s["town"], s["closed_on"]) == ("MpKrakAlKras", "Kraków", None)
    assert s["lat"] == pytest.approx(50.057678) and s["opened_on"] == dt.date(2003, 1, 1)
    assert ctx.storage.read("bronze", "gios_positions").first()["pollutant"] == "PM10"
